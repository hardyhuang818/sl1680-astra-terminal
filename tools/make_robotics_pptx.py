#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成两份 PPTX：
   1) Synaptics 博客中文版（按原文结构逐节）
   2) 代理培训版（话术/决策树/异议应答/实证）
基于 Synaptics 官方博客 "Why High-Speed Interfaces Are Essential for Modern Robotics"
(Harsha Rao, 2026-07-23)，版权归 Synaptics 所有。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = r"D:\Claude code\Case6_Astra"

NAVY   = RGBColor(0x07, 0x20, 0x4A)
BLUE   = RGBColor(0x0B, 0x6E, 0xFD)
TEAL   = RGBColor(0x00, 0xA3, 0xA3)
INK    = RGBColor(0x19, 0x1C, 0x22)
SUB    = RGBColor(0x59, 0x61, 0x6E)
LINE   = RGBColor(0xE2, 0xE6, 0xEE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT  = RGBColor(0xFA, 0xFB, 0xFD)
GOOD   = RGBColor(0x0F, 0x7B, 0x40)
GOODBG = RGBColor(0xEA, 0xF7, 0xF0)
BAD    = RGBColor(0xB9, 0x1C, 0x1C)
BADBG  = RGBColor(0xFD, 0xEC, 0xEB)
WARNBG = RGBColor(0xFF, 0xF8, 0xE6)
WARNFG = RGBColor(0x5A, 0x4A, 0x20)

FONT = "Microsoft YaHei"
W, H = Inches(13.333), Inches(7.5)


def new_deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.05
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def _tb(slide, x, y, w, h, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return tf


def _para(tf, text, size=16, color=INK, bold=False, space_after=6,
          align=PP_ALIGN.LEFT, first=False, italic=False, space_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    runs = text.split("**")
    for i, seg in enumerate(runs):
        if not seg:
            continue
        r = p.add_run(); r.text = seg
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italic
    return p


def title_slide(prs, kicker, title, subtitle, meta):
    s = _blank(prs)
    bg = _rect(s, 0, 0, W, H, fill=NAVY)
    bg.adjustments[0] = 0.0
    band = _rect(s, 0, Inches(4.6), W, Inches(0.14), fill=TEAL)
    band.adjustments[0] = 0.0
    tf = _tb(s, Inches(0.95), Inches(1.75), Inches(11.5), Inches(3.4))
    _para(tf, kicker, size=15, color=RGBColor(0x9F, 0xC6, 0xF5), first=True, space_after=12)
    _para(tf, title, size=40, color=WHITE, bold=True, space_after=10)
    _para(tf, subtitle, size=18, color=RGBColor(0xD5, 0xE4, 0xF7), space_after=4)
    tf2 = _tb(s, Inches(0.95), Inches(5.05), Inches(11.5), Inches(1.4))
    _para(tf2, meta, size=13, color=RGBColor(0xA9, 0xBE, 0xDA), first=True)
    return s


def section_slide(prs, num, title, sub=""):
    s = _blank(prs)
    bg = _rect(s, 0, 0, W, H, fill=NAVY); bg.adjustments[0] = 0.0
    tf = _tb(s, Inches(1.1), Inches(2.7), Inches(11.2), Inches(2.2))
    _para(tf, num, size=54, color=TEAL, bold=True, first=True, space_after=4)
    _para(tf, title, size=34, color=WHITE, bold=True, space_after=8)
    if sub:
        _para(tf, sub, size=16, color=RGBColor(0xB9, 0xCD, 0xE6))
    return s


def content_slide(prs, title, eyebrow=None):
    s = _blank(prs)
    top = _rect(s, 0, 0, W, Inches(1.02), fill=NAVY); top.adjustments[0] = 0.0
    accent = _rect(s, 0, Inches(1.02), W, Inches(0.06), fill=TEAL); accent.adjustments[0] = 0.0
    tf = _tb(s, Inches(0.6), Inches(0.14), Inches(12.2), Inches(0.85))
    if eyebrow:
        _para(tf, eyebrow, size=11.5, color=RGBColor(0x8F, 0xB8, 0xEB), first=True, space_after=1)
        _para(tf, title, size=25, color=WHITE, bold=True)
    else:
        _para(tf, title, size=25, color=WHITE, bold=True, first=True)
    return s


def bullets(slide, items, x=Inches(0.85), y=Inches(1.5), w=Inches(11.6), h=Inches(5.4),
            size=17, gap=13):
    tf = _tb(slide, x, y, w, h)
    first = True
    for it in items:
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        p = _para(tf, ("• " if lvl == 0 else "– ") + txt,
                  size=size if lvl == 0 else size - 2,
                  color=INK if lvl == 0 else SUB,
                  first=first, space_after=gap if lvl == 0 else gap - 5)
        p.level = lvl
        if lvl:
            from pptx.util import Inches as _I
            p.space_before = Pt(0)
        first = False
    return tf


def cards(slide, items, y=Inches(1.6), cols=4, cw=None, ch=Inches(1.55),
          gap=Inches(0.22), x0=Inches(0.75), title_size=14, body_size=12):
    if cw is None:
        total = W - Inches(1.5)
        cw = Emu(int((total - gap * (cols - 1)) / cols))
    for i, (t, b) in enumerate(items):
        r, c = divmod(i, cols)
        x = x0 + (cw + gap) * c
        yy = y + (ch + gap) * r
        _rect(slide, x, yy, cw, ch, fill=LIGHT, line=LINE)
        tf = _tb(slide, x + Inches(0.14), yy + Inches(0.11), cw - Inches(0.28), ch - Inches(0.22))
        _para(tf, t, size=title_size, color=BLUE, bold=True, first=True, space_after=4)
        _para(tf, b, size=body_size, color=SUB, space_after=0)


def table(slide, headers, rows, x=Inches(0.7), y=Inches(1.55),
          w=None, col_w=None, fs=12.5, hfs=13, rh=Inches(0.42)):
    if w is None:
        w = W - Inches(1.4)
    nrows, ncols = len(rows) + 1, len(headers)
    h = rh * nrows
    gf = slide.shapes.add_table(nrows, ncols, x, y, w, h)
    tbl = gf.table
    tbl.first_row = True
    if col_w:
        tot = sum(col_w)
        for i, cwv in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w * cwv / tot))
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = Inches(0.08)
        p = c.text_frame.paragraphs[0]
        r = p.add_run(); r.text = htxt
        r.font.size = Pt(hfs); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = FONT
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            p = c.text_frame.paragraphs[0]
            for k, seg in enumerate(str(val).split("**")):
                if not seg:
                    continue
                r = p.add_run(); r.text = seg
                r.font.size = Pt(fs); r.font.name = FONT
                r.font.color.rgb = INK
                r.font.bold = (k % 2 == 1)
    return tbl


def callout(slide, text, x=Inches(0.8), y=Inches(5.6), w=Inches(11.7), h=Inches(0.95),
            bg=GOODBG, fg=GOOD, label=None, size=15):
    _rect(slide, x, y, w, h, fill=bg, line=fg, line_w=1.0)
    bar = _rect(slide, x, y, Inches(0.07), h, fill=fg); bar.adjustments[0] = 0.0
    tf = _tb(slide, x + Inches(0.22), y + Inches(0.11), w - Inches(0.45), h - Inches(0.22))
    if label:
        _para(tf, label, size=11.5, color=fg, bold=True, first=True, space_after=3)
        _para(tf, text, size=size, color=INK, space_after=0)
    else:
        _para(tf, text, size=size, color=INK, bold=False, first=True, space_after=0)


def mono_block(slide, lines, x=Inches(0.8), y=Inches(1.55), w=Inches(11.7), h=Inches(5.0),
               size=12.5):
    bx = _rect(slide, x, y, w, h, fill=RGBColor(0x0D, 0x1B, 0x2E))
    tf = _tb(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), h - Inches(0.36))
    first = True
    for ln in lines:
        p = _para(tf, ln if ln else " ", size=size,
                  color=RGBColor(0xD8, 0xE4, 0xF5), first=first, space_after=2)
        for r in p.runs:
            r.font.name = "Consolas"
        first = False


def end_slide(prs, lines, note):
    s = _blank(prs)
    bg = _rect(s, 0, 0, W, H, fill=NAVY); bg.adjustments[0] = 0.0
    tf = _tb(s, Inches(1.2), Inches(2.1), Inches(11.0), Inches(3.2))
    first = True
    for ln in lines:
        _para(tf, ln, size=26, color=WHITE, bold=True, first=first, space_after=18)
        first = False
    tf2 = _tb(s, Inches(1.2), Inches(5.7), Inches(11.0), Inches(1.2))
    _para(tf2, note, size=12, color=RGBColor(0x9F, 0xB6, 0xD4), first=True)


# ══════════════════════════════════════════════════════════════════
# 第一份：中文版
# ══════════════════════════════════════════════════════════════════
def build_cn():
    prs = new_deck()

    title_slide(
        prs,
        "SYNAPTICS 官方博客 · 中文版",
        "为什么高速接口对现代机器人至关重要",
        "Why High-Speed Interfaces Are Essential for Modern Robotics",
        "原文作者：Harsha Rao，Synaptics 视频与数据接口事业部 VP & GM　|　2026-07-23　|　分类：AI、边缘计算\n"
        "原文：synaptics.com/company/blog/why-high-speed-interfaces-are-essential-for-modern-robotics\n"
        "本中文版按原文结构整理、以中文重述，版权归 Synaptics Incorporated 所有　|　整理：2026-07-28")

    s = content_slide(prs, "开篇：被忽略的那一层", "INTRODUCTION")
    bullets(s, [
        "“物理 AI”（Physical AI）是一类新范式——系统能够**感知、推理，并在物理世界中行动**",
        "机器人是物理 AI 最直观的载体",
        "行业注意力几乎全在 **AI 算力**和**传感**上",
        "而**高分辨率视频与显示数据在机器人体内怎么搬运**，被系统性忽略",
    ], y=Inches(1.75), size=19, gap=20)
    callout(s, "恰恰是这一层，决定了一台机器人能否做成、做轻、做可靠。",
            y=Inches(5.5), label="原文的核心提醒", size=17)

    s = content_slide(prs, "视觉沟通正在成为人机界面", "第 1 节 · Visual Communication Is Becoming the Human-Robot Interface")
    bullets(s, [
        "**显示屏是建立信任的主界面**——人要理解机器人在想什么、要做什么，最直接的通道就是屏幕",
        "**视觉反馈内容远不止“状态灯”**：导航信息、运行状态、诊断信息、远程视频回传、沉浸式可视化",
        "**人形机器人还要用屏幕做情绪表达**，让交互更直观、更有参与感",
    ], y=Inches(1.8), size=18, gap=22)
    callout(s, "视觉沟通已经是核心用户体验，不再是“锦上添花”的附属功能。",
            y=Inches(5.35), label="结论", size=17)

    s = content_slide(prs, "机器人正在重新定义系统架构", "第 2 节 · Robotics Is Redefining System Architecture")
    tf = _tb(s, Inches(0.85), Inches(1.4), Inches(11.6), Inches(1.0))
    _para(tf, "消费电子里处理器和屏幕紧挨着；机器人不是这样——**算力集中在躯干/底盘，显示分散在头部、手臂、操作台**，"
              "中间还要过关节、过窄颈。这直接把八个彼此冲突的约束摆上台面：",
          size=15.5, color=INK, first=True)
    cards(s, [
        ("机械复杂度", "线缆穿过转动关节，走线空间常只有几毫米"),
        ("线缆重量", "每根线都是负载，吃掉负重与续航预算"),
        ("功耗", "长距离高速传输的驱动功耗不可忽略"),
        ("信号完整性", "长线、过弯、活动件旁边，眼图会塌"),
        ("散热", "接口芯片与驱动器的热量要有去处"),
        ("电磁兼容", "高速线与电机驱动器共处一体，EMC 是硬骨头"),
        ("可靠性", "活动线缆是整机故障率最高的部位之一"),
        ("制造成本", "线束是装配工时和不良率的主要来源"),
    ], y=Inches(2.55), cols=4, ch=Inches(1.42))
    callout(s, "减少线缆不是省钱问题，而是**直接改善系统功能**。",
            y=Inches(6.05), h=Inches(0.78), label="作者的核心判断", size=16)

    s = content_slide(prs, "高速接口在机器人系统里到底做什么", "第 3 节 · What High-Speed Interfaces Do")
    tf = _tb(s, Inches(0.85), Inches(1.4), Inches(11.6), Inches(0.9))
    _para(tf, "**连接对象：**摄像头与传感器 · 头部显示 · 操作员 HMI · 维修/服务面板 · 外围显示 · 远程遥操作系统",
          size=15.5, color=INK, first=True)
    table(s, ["目标（不只是把像素送过去）", "含义"], [
        ["降低布线复杂度", "用更少的线、更细的线走完同样的路径"],
        ["保持信号完整性", "长距离、过关节之后画面依然干净"],
        ["压低延迟", "遥操作和精密可视化对延迟极度敏感"],
        ["提升功效比", "同样带宽下更省电，续航更长"],
        ["架构灵活性", "系统怎么摆由需求决定，而不是被接口卡住"],
    ], y=Inches(2.25), col_w=[3, 7], rh=Inches(0.62))

    s = content_slide(prs, "为什么机器人需要“可选择”的显示接口", "第 4 节 · Why Robotics Requires Flexible Display Interfaces")
    bullets(s, [
        "**AI 处理器的显示能力差异极大**——有的自带 GPU 和原生 DisplayPort 输出；有的把晶体管全押在加速器上，显示能力很弱甚至没有",
        "**因此必须做协议转换**，在 DisplayPort、MIPI、LVDS、USB 以及车载视频接口之间自由搭桥",
        "**灵活的接口让工程师按整机最优来设计**，而不是被某颗主控的接口限制反向绑架整个架构",
    ], y=Inches(1.8), size=18, gap=22)
    callout(s, "当主控显示能力有限时，**基于 USB 或以太网的视频传输是一条高效的替代路径**——这正是 DisplayLink 的定位。",
            y=Inches(5.35), label="原文特别提到", size=16)

    s = content_slide(prs, "从架构到落地：三个真实案例", "第 5 节 · From Architecture to Deployment")
    table(s, ["案例", "痛点", "方案", "收益"], [
        ["人形机器人\n（头显）", "主控原生显示能力有限，头部要驱动显示屏",
         "**DisplayLink® 基于 USB 的视频传输架构**", "大幅减少布线与机械复杂度"],
        ["人形机器人\n（窄颈多屏）", "多路高分辨率显示，线缆必须穿过狭窄的颈部组件",
         "**长距离视频传输 + 协议桥接**到嵌入式显示接口", "在极其受限的机械空间内支持多路高分屏"],
        ["医疗机器人\n平台", "精密可视化要求**多路同步 4K**，延迟极低",
         "**高带宽 DisplayPort 架构**，传输**未压缩**视频流", "满足医疗场景对画质与实时性的双重刚需"],
    ], y=Inches(1.7), col_w=[1.5, 3.2, 3.4, 2.9], rh=Inches(1.25), fs=13)

    s = content_slide(prs, "为什么“减线”能改善整机设计", "第 6 节 · Why Cable Reduction Improves Robotic System Design")
    cards(s, [
        ("运动自由度", "线少了，关节活动范围和灵活性都能放开"),
        ("机械简化", "走线通道、拖链、护套统统简化"),
        ("可制造性", "线束少 → 装配工时短、错装概率低"),
        ("可维修性", "拆装更快，售后成本下降"),
        ("电池效率", "轻量化 + 低功耗接口 = 续航变长"),
        ("长期可靠性", "活动线缆是疲劳失效高发点，减少即降故障率"),
    ], y=Inches(1.75), cols=3, ch=Inches(1.6))
    callout(s, "选型时要看的不只是带宽，还包括：**延迟、信号完整性、功效、安全性、可扩展性**。",
            y=Inches(5.5), label="原文强调", size=16)

    s = content_slide(prs, "Synaptics 的视角与产品线", "第 7 节 · The Synaptics Perspective")
    table(s, ["产品线", "做什么", "典型用法"], [
        ["**DisplayLink 解码器**",
         "把视频通过 **USB 或以太网**送出去，末端解码后驱动嵌入式 **eDP / MIPI / LVDS** 屏",
         "主控显示能力弱、或需要长距离/单线缆传输"],
        ["**DisplayPort 桥接 IC**",
         "把 **DisplayPort** 转成嵌入式显示接口：**eDP / MIPI DSI / LVDS**",
         "主控有原生 DP，但屏是嵌入式接口"],
        ["**DisplayPort MST Hub**",
         "用**一条高带宽 DP** 驱动**多路独立显示**",
         "多屏系统，且要把线缆数量压到最低"],
    ], y=Inches(1.7), col_w=[2.3, 5.3, 3.4], rh=Inches(1.15), fs=13)
    callout(s, "三条产品线的共同价值：**减少布线 · 简化机械走线 · 支持长距离视频传输**",
            y=Inches(5.55), label="共同价值主张", size=16)

    s = content_slide(prs, "展望", "第 8 节 · Looking Ahead")
    bullets(s, [
        "机器人能否成功，取决于**完整的系统架构**——信息能否在**感知 → 计算 → 人机交互**之间高效流动",
        "视觉沟通将定义用户体验：**富有表情的人形机器人面部**、**沉浸式手术显示**、**工业 HMI**、**远程遥操作台**",
        "**高速接口是机器人技术栈中的一个战略层**，不是可有可无的配件",
    ], y=Inches(1.85), size=18, gap=24)
    callout(s, "Synaptics 的定位：把数十年的连接技术积累用于帮助开发者解决这些问题，"
               "让高速接口成为未来机器人设计与部署的基础能力。",
            y=Inches(5.3), h=Inches(1.0), label="结语", size=15.5)

    end_slide(prs, [
        "减少线缆不是省钱问题，是直接改善系统功能。",
        "高速接口是机器人技术栈里的一个战略层。",
    ], "原文版权归 Synaptics Incorporated 所有。本中文版按原文结构整理、以中文重述，供内部阅读与培训参考；\n"
       "对外引用请注明来源并附原文链接。整理日期：2026-07-28")

    path = os.path.join(OUT, "Synaptics博客_机器人为何需要高速接口_中文_2026-07-28.pptx")
    prs.save(path)
    return path, len(prs.slides.__iter__.__self__._sldIdLst)


# ══════════════════════════════════════════════════════════════════
# 第二份：代理培训版
# ══════════════════════════════════════════════════════════════════
def build_train():
    prs = new_deck()

    title_slide(
        prs,
        "内部培训材料 · 代理商版",
        "机器人高速接口：怎么讲、怎么选、怎么答",
        "销售与售前培训包 · 建议 45 分钟",
        "基于 Synaptics 官方博客《Why High-Speed Interfaces Are Essential for Modern Robotics》\n"
        "（Harsha Rao，2026-07-23）整理；话术、决策树、异议应答与实证部分为本团队补充\n"
        "整理：2026-07-28　|　本材料供内部培训使用")

    s = content_slide(prs, "培训大纲", "AGENDA · 45 min")
    left = ["1　一句话卖点：先把话说对", "2　市场为什么现在才需要它",
            "3　客户痛点地图：八个约束", "4　三条产品线各自解决什么", "5　选型决策树"]
    right = ["6　三个真实案例怎么讲", "7　客户画像与开场话术",
             "8　常见异议与应答", "9　我们自己的实证弹药", "10　一页速查表"]
    for col, items, x in ((0, left, Inches(1.0)), (1, right, Inches(7.0))):
        tf = _tb(s, x, Inches(1.9), Inches(5.4), Inches(4.6))
        first = True
        for it in items:
            _para(tf, it, size=19, color=INK, first=first, space_after=22)
            first = False

    section_slide(prs, "01", "一句话卖点", "先把话说对")
    s = content_slide(prs, "一句话卖点：先把话说对", "第 1 节")
    callout(s, "“机器人的算力在躯干，屏幕在头上、手上、操作台上——中间那段怎么走线，"
               "决定了这台机器人能不能做轻、做可靠、做得出来。\n"
               "**我们卖的不是一颗接口芯片，是帮你少走十几根线。**”",
            y=Inches(1.6), h=Inches(1.6), label="✔ 这样讲", size=17)
    callout(s, "“我们有 DisplayLink、DP Bridge、MST Hub 三个产品线，带宽分别是……”\n"
               "上来讲料号和参数，客户听不出跟他有什么关系。**先讲他的机械和续航痛点。**",
            y=Inches(3.5), h=Inches(1.35), bg=BADBG, fg=BAD, label="✘ 别这样讲", size=16)
    callout(s, "原文最有力的判断是——**“减少线缆不是省钱问题，而是直接改善系统功能”**。"
               "这句话是整个销售逻辑的支点，务必背下来并会展开。",
            y=Inches(5.25), h=Inches(1.0), bg=WARNBG, fg=RGBColor(0xB4, 0x53, 0x09),
            label="培训要点", size=15.5)

    s = content_slide(prs, "市场为什么现在才需要它", "第 2 节")
    bullets(s, [
        "**“物理 AI”成为新范式**：系统要能感知、推理，并在物理世界中行动——机器人是最直观的载体",
        "**行业注意力全在算力和传感上**，视频/显示数据在机体内怎么搬运被系统性忽略 → **这就是我们的机会窗口**",
        "**屏幕从“状态灯”升级成“主界面”**：导航、状态、诊断、远程视频回传、沉浸式可视化，还有人形机器人的**表情表达**",
    ], y=Inches(1.8), size=17.5, gap=22)
    callout(s, "**给客户的结论：**视觉沟通已经是核心用户体验，不是附属功能——所以显示链路值得单独做架构设计。",
            y=Inches(5.4), label="讲法", size=16)

    s = content_slide(prs, "客户痛点地图：八个约束", "第 3 节")
    tf = _tb(s, Inches(0.85), Inches(1.35), Inches(11.6), Inches(0.6))
    _para(tf, "现场提问时，从这八个里**挑客户最疼的两个**切入：", size=15.5, color=INK, first=True)
    cards(s, [
        ("机械复杂度", "线要穿转动关节，走线空间常只有几毫米"),
        ("线缆重量", "每根线都是负载，直接吃负重与续航预算"),
        ("功耗", "长距离高速传输的驱动功耗不可忽略"),
        ("信号完整性", "长线、过弯、贴着电机走，眼图会塌"),
        ("散热", "接口芯片与驱动器的热量要有去处"),
        ("电磁兼容", "高速线与电机驱动器同处一体，EMC 是硬骨头"),
        ("可靠性", "活动线缆是整机故障率最高的部位之一"),
        ("制造成本", "线束是装配工时与不良率的主要来源"),
    ], y=Inches(2.0), cols=4, ch=Inches(1.42))
    callout(s, "“您这台机器人，**颈部转动范围内要过几根线**？做疲劳测试时**线缆是不是第一个坏的**？”\n"
               "这两个问题几乎必中，问完客户自己就把痛点讲出来了。",
            y=Inches(5.5), h=Inches(1.15), label="✔ 这样问", size=16)

    s = content_slide(prs, "三条产品线各自解决什么", "第 4 节")
    table(s, ["产品线", "技术上做什么", "什么时候推它", "客户获得的收益"], [
        ["**DisplayLink\n解码器**",
         "视频经 **USB 或以太网**传输，末端解码驱动嵌入式 **eDP / MIPI / LVDS** 屏",
         "主控**没有**或只有很弱的显示输出；需要长距离、单线缆",
         "一根线走完视频+数据+供电，布线最省"],
        ["**DisplayPort\n桥接 IC**",
         "把 **DisplayPort** 转成 **eDP / MIPI DSI / LVDS**",
         "主控**有**原生 DP，但屏是嵌入式接口",
         "不改主控、不改屏，中间加一颗桥片就打通"],
        ["**DisplayPort\nMST Hub**",
         "一条高带宽 **DP** 驱动**多路独立显示**",
         "多屏系统，且线缆数量必须压到最低",
         "N 块屏只走一条线，关节处线束大幅缩减"],
    ], y=Inches(1.6), col_w=[1.7, 3.6, 3.2, 2.9], rh=Inches(1.2), fs=12.5)
    callout(s, "**DisplayLink** 解决“主控给不出视频”　|　**DP Bridge** 解决“格式对不上”　|　"
               "**MST Hub** 解决“屏太多线太多”　—— 三者常组合出现在同一台机器上",
            y=Inches(5.55), h=Inches(0.95), bg=WARNBG, fg=RGBColor(0xB4, 0x53, 0x09),
            label="关键区分", size=15.5)

    s = content_slide(prs, "选型决策树（现场就能用）", "第 5 节")
    mono_block(s, [
        "客户的 AI 主控有原生视频输出吗？",
        "│",
        "├─ 没有 / 很弱 ─────────────────────► DisplayLink（USB 或以太网视频传输）",
        "│                                      末端解码 → eDP / MIPI / LVDS",
        "│                                      卖点：一根线走完，布线最省",
        "│",
        "└─ 有原生 DisplayPort",
        "   │",
        "   ├─ 只驱动 1 块屏，屏是嵌入式接口 ─► DP Bridge IC",
        "   │                                   DP → eDP / MIPI DSI / LVDS",
        "   │                                   卖点：不动主控、不动屏",
        "   │",
        "   └─ 要驱动 2 块以上独立显示",
        "      │",
        "      ├─ 线缆数量是硬约束（过关节）─► MST Hub（末端按需加 Bridge）",
        "      │                               卖点：一条 DP 出去，多屏落地",
        "      │",
        "      └─ 多路同步 4K + 极低延迟 ────► 高带宽 DisplayPort，未压缩传输",
        "                                      卖点：医疗/精密可视化刚需",
    ], y=Inches(1.5), h=Inches(4.55), size=12.5)
    callout(s, "**必问四问：**①主控型号和视频输出？②屏的接口、分辨率、几块？③线走多长、过不过活动关节？④延迟和压缩能不能接受？",
            y=Inches(6.2), h=Inches(0.8), bg=WARNBG, fg=RGBColor(0xB4, 0x53, 0x09), size=15)

    s = content_slide(prs, "三个真实案例怎么讲", "第 6 节")
    table(s, ["案例", "客户痛点", "采用方案", "收益", "打什么客户"], [
        ["**人形机器人**\n（头显）", "主控原生显示能力有限，头部要驱动屏",
         "**DisplayLink® USB 视频传输架构**", "显著减少布线与机械复杂度", "人形本体厂\n服务机器人"],
        ["**人形机器人**\n（窄颈多屏）", "多路高分屏，线必须穿过狭窄颈部",
         "**长距离视频传输 + 协议桥接**", "极限机械空间内支持多路高分屏", "结构受限严重的\n整机厂"],
        ["**医疗机器人**\n平台", "多路**同步 4K**，延迟要求极低",
         "**高带宽 DP 架构，未压缩流**", "满足精密可视化画质+实时性", "手术机器人\n影像/精密检测"],
    ], y=Inches(1.6), col_w=[1.6, 2.6, 2.9, 2.7, 1.9], rh=Inches(1.2), fs=12)
    callout(s, "按“**痛点 → 方案 → 量化收益**”三段式讲，不要念产品参数。"
               "客户记不住带宽数字，但记得住“**穿过脖子的线从十几根变成一根**”。",
            y=Inches(5.55), h=Inches(0.95), label="✔ 讲法", size=15.5)

    s = content_slide(prs, "客户画像与开场话术", "第 7 节")
    table(s, ["客户类型", "他最在意", "开场切入"], [
        ["人形/服务机器人本体厂", "负重、续航、关节走线、量产良率", "“头部这几块屏，线是怎么过脖子的？”"],
        ["医疗 / 手术机器人", "延迟、画质、认证、可靠性", "“4K 是几路同步？能接受压缩吗？”"],
        ["工业 HMI / AGV", "EMC、维修性、长距离", "“控制柜到操作面板多远？现在走什么线？”"],
        ["遥操作 / 巡检", "长距传输、端到端延迟", "“远端画面回来延迟多少？操作员抱怨吗？”"],
        ["AI 主控 / 模组商", "自家芯片显示能力短板", "“客户要多屏时，你们怎么补显示输出？”"],
    ], y=Inches(1.6), col_w=[3, 3.6, 4.6], rh=Inches(0.72), fs=13)
    callout(s, "**共同抓手：**不管哪一类，**先问“线”，别先问“屏”**。线是他们真正头疼、且说得出数字的东西。",
            y=Inches(5.6), bg=WARNBG, fg=RGBColor(0xB4, 0x53, 0x09), size=16)

    section_slide(prs, "08", "常见异议与应答", "五个最常碰到的问题")

    qas = [
        ("Q1　我主控自带 HDMI/DP，直接拉线不就行了？",
         "能通不等于能量产。反问三件事：这根线要不要**过活动关节**？做完 **10 万次弯折疲劳**测试还通不通？**EMC** 过不过？\n"
         "直连方案在实验室能跑，在整机上是故障率和装配工时的来源。而且主控只有一路输出，客户后面**加第二块屏就得推倒重来**。"),
        ("Q2　USB 传视频会不会有延迟/压缩，影响体验？",
         "分场景。表情屏、状态屏、HMI 这类，**USB 传输完全够用**，收益是布线大幅简化。\n"
         "但如果是**多路同步 4K、精密可视化**（比如手术），就该走**高带宽 DisplayPort、传未压缩流**——原文的医疗案例就是这么做的。\n"
         "**我们不是只有一条路，是按场景选架构。**"),
        ("Q3　加一颗桥片，成本和板面积都上去了。",
         "把账算全。桥片省下来的是：**线束物料、装配工时、走线机械件、疲劳失效的售后成本**，以及**重量换来的续航**。\n"
         "原文的原话是——减少线缆的收益覆盖**运动自由度、机械简化、可制造性、可维修性、电池效率、长期可靠性**，"
         "**这些都不在 BOM 那一栏里**。"),
        ("Q4　我们已经在用某家的 LVDS/MIPI 桥了。",
         "那说明架构已经定型在“主控出什么、屏收什么”这条线上。反问一句：**下一代要加屏的时候怎么办？**\n"
         "MST Hub 和 DisplayLink 的价值是让**架构可扩展**，而不是每加一块屏就重新设计线束。\n"
         "另外做桥片选型时，**带宽余量**是最容易被忽略的坑（见下一节实证）。"),
        ("Q5　这跟车载那套是一回事吗？",
         "思路一致，接口不同。原文明确提到协议转换要覆盖 **DisplayPort、MIPI、LVDS、USB 以及车载视频接口**。\n"
         "机器人的特殊之处是**活动关节**和**电池供电**，所以对**重量和弯折寿命**比车载更敏感。"),
    ]
    for i in range(0, len(qas), 2):
        s = content_slide(prs, "常见异议与应答", "第 8 节 · %d/3" % (i // 2 + 1))
        y = Inches(1.55)
        for q, a in qas[i:i + 2]:
            _rect(s, Inches(0.8), y, Inches(11.7), Inches(2.35), fill=LIGHT, line=LINE)
            bar = _rect(s, Inches(0.8), y, Inches(0.07), Inches(2.35), fill=BLUE)
            bar.adjustments[0] = 0.0
            tf = _tb(s, Inches(1.05), y + Inches(0.15), Inches(11.2), Inches(2.05))
            _para(tf, q, size=16.5, color=NAVY, bold=True, first=True, space_after=8)
            for ln in a.split("\n"):
                _para(tf, ln, size=14, color=INK, space_after=4)
            y = y + Inches(2.6)

    section_slide(prs, "09", "我们自己的实证弹药", "讲给客户听，比念 datasheet 可信得多")

    s = content_slide(prs, "实证一 · 多屏确实靠一条线做到了", "第 9 节 · SL1680（Synaptics Astra）项目一手经验")
    bullets(s, [
        "用 **DisplayLink（DL7400）**在 SL1680 上实测跑通**三屏**输出，主控本身并不具备三路原生显示能力",
        "**踩过的坑（讲出来更可信）：**",
        ("不要相信 EDID 报的 preferred_mode", 1),
        ("显示器断电后链路不会自愈，要在系统层做重枚举", 1),
    ], y=Inches(1.8), size=18, gap=18)
    callout(s, "**可直接对客户讲的一句话：**“主控给不出三路视频，我们用 DisplayLink 一样做出来了。”",
            y=Inches(5.4), size=16.5)

    s = content_slide(prs, "实证二 · 桥接芯片的带宽余量是生死线", "第 9 节 · TC358775（DSI→LVDS）实测")
    bullets(s, [
        "同一项目用 **TC358775** 驱动 10.5″ LVDS 屏",
        "按厂商示例配置**只留 4.4% 带宽余量** → 实测出现**整屏跑动的雪花噪点**（行缓冲欠载）",
        "把余量提到 **33%** 后画面才彻底干净",
    ], y=Inches(1.85), size=18.5, gap=22)
    callout(s, "**对客户的价值：**“选桥片别只看它标称支持多少分辨率，要看你**实际配置下还剩多少余量**。”\n"
               "这是能立刻建立专业信任的一句话。",
            y=Inches(5.2), h=Inches(1.05), size=16)

    s = content_slide(prs, "实证三 ·“线”真的是故障率最高的地方", "第 9 节")
    bullets(s, [
        "本项目的 LVDS 目前仍是**杜邦线**，是已知的信号质量风险项，计划换双绞线",
        "USB 摄像头在总线抖动时会**掉线数十秒**，必须在软件层做守护重连",
    ], y=Inches(1.9), size=18.5, gap=24)
    callout(s, "**讲法：**“我们自己在实验台上都被线缆坑过，你们装在**会动的机器人**上，只会更严重。”",
            y=Inches(4.4), size=16.5)
    callout(s, "**使用提醒：**对外讲时说“我们在 SL1680 平台的实测经验”即可，不要透露客户名称与未公开的型号细节；"
               "DisplayLink SDK 相关二进制与厂商资料受 **NDA** 约束，不得外发。",
            y=Inches(5.7), h=Inches(1.0), bg=BADBG, fg=BAD, size=15)

    s = content_slide(prs, "一页速查表", "第 10 节 · 打印出来带在身上")
    table(s, ["场景", "推什么", "一句话卖点"], [
        ["主控没有视频输出", "**DisplayLink 解码器**", "一根 USB/网线走完视频"],
        ["主控有 DP，屏是 MIPI/LVDS/eDP", "**DP Bridge IC**", "不动主控、不动屏，中间搭个桥"],
        ["多屏 + 线缆数量硬约束", "**MST Hub**", "一条 DP 出去，多屏落地"],
        ["多路同步 4K + 极低延迟", "**高带宽 DP，未压缩**", "医疗/精密可视化的唯一解"],
        ["长距离到操作台", "**DisplayLink（USB/以太网）**", "长距传输不牺牲布线简洁"],
    ], y=Inches(1.55), col_w=[3.6, 3.4, 4.2], rh=Inches(0.7), fs=13)
    callout(s, "**必问四问：**①主控型号/视频输出　②屏接口·分辨率·数量　③线长·过不过活动关节　④延迟·能否压缩",
            y=Inches(5.5), h=Inches(0.8), bg=WARNBG, fg=RGBColor(0xB4, 0x53, 0x09), size=15)

    end_slide(prs, [
        "① 减少线缆不是省钱问题，是直接改善系统功能。",
        "② 高速接口是机器人技术栈里的一个战略层。",
        "③ 让架构按整机最优设计，别被主控接口绑架。",
    ], "核心观点、案例与产品定位来自 Synaptics 官方博客《Why High-Speed Interfaces Are Essential for Modern Robotics》\n"
       "（Harsha Rao，2026-07-23），版权归 Synaptics Incorporated 所有。销售话术、决策树、异议应答与实证内容为本团队补充。\n"
       "本材料供内部培训使用；对外分发前请确认 NDA 相关内容已移除。整理：2026-07-28")

    path = os.path.join(OUT, "Synaptics博客_代理培训_机器人高速接口_2026-07-28.pptx")
    prs.save(path)
    return path, len(prs.slides._sldIdLst)


if __name__ == "__main__":
    p1 = build_cn()
    p2 = build_train()
    for p in (p1, p2):
        print("  %s  (%d 页, %.0f KB)" % (os.path.basename(p[0]), p[1],
                                          os.path.getsize(p[0]) / 1024))
